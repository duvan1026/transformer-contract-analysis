"""
Escritura y formateo del Excel de salida del análisis de longitud de tokens.
Responsabilidad separada de token_length_stats.py: aquí solo vive la
presentación (formato de celdas, anchos de columna, hoja de diccionario),
no el cálculo.

Incluye una hoja "00_Diccionario_Variables" generada automáticamente en
cada corrida, para que el archivo quede autoexplicativo sin depender de
documentación externa que se pueda desincronizar.
"""

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# Columnas que se muestran como porcentaje en el formato de celda de Excel
_COLUMNAS_PORCENTAJE = {
    "Porcentaje_Excede_Limite_Tokens",
    "Porcentaje_En_Rango_Riesgo_Tokens",
    "Porcentaje_Mayor_Umbral_Palabras",
}

# Color de encabezado por bloque temático (mismo criterio usado al explicar las variables)
_COLOR_GENERAL = "D9D9D9"   # Clase, Cantidad_Fragmentos
_COLOR_TOKENS = "C6D9F1"    # columnas Tokens_*
_COLOR_PALABRAS = "D8E4BC"  # columnas Palabras_*


def _construir_diccionario_variables(umbral_palabras, zona_riesgo_tokens, max_tokens):
    """
    Construye la hoja de diccionario de variables. Los valores de umbral se
    interpolan dinámicamente desde text_standards.py, así que si cambian,
    la descripción en el Excel queda actualizada automáticamente.
    """
    filas = [
        ("Clase", "Categoría de cláusula evaluada.", "Texto",
         "Agrupa todas las demás métricas de la fila."),
        ("Cantidad_Fragmentos", "Número total de fragmentos/spans de esa clase en esa fuente.", "Conteo",
         "A mayor cantidad, más confiables son los percentiles calculados sobre esa fila."),
        ("Tokens_Mediana", "Valor de tokens que deja la mitad de los fragmentos por debajo y la mitad por arriba.",
         "Tokens", "Es el tamaño 'típico' del fragmento en esa clase y fuente."),
        ("Tokens_Percentil_75", "Valor de tokens por debajo del cual cae el 75% de los fragmentos.", "Tokens",
         "Empieza a mostrar la cola de fragmentos más largos de lo normal."),
        ("Tokens_Percentil_90", "Valor de tokens por debajo del cual cae el 90% de los fragmentos.", "Tokens",
         "Muestra la cola larga con más claridad que el percentil 75."),
        ("Tokens_Percentil_95", "Valor de tokens por debajo del cual cae el 95% de los fragmentos.", "Tokens",
         "Cola larga extrema; útil para detectar clases con fragmentos inusualmente largos."),
        ("Tokens_Maximo", "Longitud en tokens del fragmento más largo observado.", "Tokens",
         "Detecta outliers extremos que los percentiles pueden no capturar."),
        ("Porcentaje_Excede_Limite_Tokens",
         f"Porcentaje de fragmentos que superan el límite físico de tokens del modelo (actualmente {max_tokens}).",
         "%", "Solo aplica a train/val/test. Un valor mayor a 0 indica truncamiento silencioso al entrenar."),
        ("Cantidad_Excede_Limite_Tokens", "Número absoluto de fragmentos que superan ese límite.", "Conteo",
         "Complementa la columna de porcentaje con el valor bruto de fragmentos afectados."),
        ("Porcentaje_En_Rango_Riesgo_Tokens",
         f"Porcentaje de fragmentos cuya longitud cae en el rango de tokens "
         f"({zona_riesgo_tokens[0]}-{zona_riesgo_tokens[1]}) donde se observó menor acierto del modelo.", "%",
         "Cuantifica qué proporción del volumen real cae en una zona de riesgo ya identificada. "
         "Este rango es configurable (TOKEN_LENGTH_RISK_ZONE_TOKENS en text_standards.py) y puede cambiar."),
        ("Cantidad_En_Rango_Riesgo_Tokens", "Número absoluto de fragmentos en ese rango.", "Conteo",
         "Complementa la columna de porcentaje con el valor bruto de fragmentos afectados."),
        ("Palabras_Mediana", "Igual que Tokens_Mediana, pero contando palabras separadas por espacios.", "Palabras",
         "Permite comparar contra criterios de longitud definidos en palabras, no en tokens del modelo."),
        ("Palabras_Percentil_75", "Igual que Tokens_Percentil_75, en palabras.", "Palabras", "Ver Tokens_Percentil_75."),
        ("Palabras_Percentil_90", "Igual que Tokens_Percentil_90, en palabras.", "Palabras", "Ver Tokens_Percentil_90."),
        ("Palabras_Percentil_95", "Igual que Tokens_Percentil_95, en palabras.", "Palabras", "Ver Tokens_Percentil_95."),
        ("Palabras_Maximo", "Longitud en palabras del fragmento más largo observado.", "Palabras",
         "Ver Tokens_Maximo."),
        ("Porcentaje_Mayor_Umbral_Palabras",
         f"Porcentaje de fragmentos que superan las {umbral_palabras} palabras.", "%",
         "Umbral configurable (TOKEN_LENGTH_WORD_THRESHOLD en text_standards.py). "
         "Sirve para detectar fragmentos anómalamente largos independientemente del conteo de tokens."),
        ("Cantidad_Mayor_Umbral_Palabras", "Número absoluto de fragmentos que superan ese umbral de palabras.",
         "Conteo", "Complementa la columna de porcentaje con el valor bruto de fragmentos afectados."),
    ]
    return pd.DataFrame(filas, columns=["Variable", "Que_Mide", "Unidad", "Como_Interpretarla"])


def _autoajustar_columnas(worksheet):
    """Ajusta el ancho de cada columna al contenido más largo (encabezado o valor)."""
    for col_cells in worksheet.columns:
        longitud_max = 0
        col_letra = None
        for cell in col_cells:
            if col_letra is None:
                col_letra = get_column_letter(cell.column)
            valor = "" if cell.value is None else str(cell.value)
            longitud_max = max(longitud_max, len(valor))
        if col_letra:
            worksheet.column_dimensions[col_letra].width = min(longitud_max + 3, 60)


def _formatear_hoja_estadisticos(worksheet, df):
    """Aplica encabezado con color por bloque, formato %, freeze panes y ancho automático."""
    header_font = Font(bold=True, color="000000")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, col_name in enumerate(df.columns, start=1):
        cell = worksheet.cell(row=1, column=idx)
        cell.font = header_font
        cell.alignment = header_alignment

        if col_name in ("Clase", "Cantidad_Fragmentos"):
            color = _COLOR_GENERAL
        elif col_name.startswith("Tokens_") or "Tokens" in col_name:
            color = _COLOR_TOKENS
        elif col_name.startswith("Palabras_") or "Palabras" in col_name:
            color = _COLOR_PALABRAS
        else:
            color = _COLOR_GENERAL
        cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

    # Formato de porcentaje real en las columnas correspondientes
    for idx, col_name in enumerate(df.columns, start=1):
        if col_name in _COLUMNAS_PORCENTAJE:
            for row in range(2, len(df) + 2):
                worksheet.cell(row=row, column=idx).number_format = '0.0"%"'

    worksheet.freeze_panes = "A2"
    _autoajustar_columnas(worksheet)


def _formatear_hoja_diccionario(worksheet, df):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    wrap_alignment = Alignment(vertical="top", wrap_text=True)

    for idx in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, len(df) + 2):
        for col in range(1, len(df.columns) + 1):
            worksheet.cell(row=row, column=col).alignment = wrap_alignment

    worksheet.freeze_panes = "A2"
    # Anchos fijos más generosos para las columnas de texto largo
    anchos = {"A": 32, "B": 55, "C": 12, "D": 55}
    for col_letra, ancho in anchos.items():
        worksheet.column_dimensions[col_letra].width = ancho


def guardar_excel_estadisticos(stats_por_fuente, resumen_comparativo, output_path,
                                umbral_palabras, zona_riesgo_tokens, max_tokens):
    """
    Guarda el Excel completo: hoja de diccionario primero, luego una hoja
    por fuente (train/val/test/inferencia), y por último el resumen
    comparativo. Aplica formato legible a cada hoja.
    """
    df_diccionario = _construir_diccionario_variables(umbral_palabras, zona_riesgo_tokens, max_tokens)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_diccionario.to_excel(writer, sheet_name="00_Diccionario_Variables", index=False)

        for fuente, df in stats_por_fuente.items():
            hoja = fuente if not df.empty else f"{fuente}_vacio"
            (df if not df.empty else pd.DataFrame({"Aviso": ["Sin datos disponibles"]})).to_excel(
                writer, sheet_name=hoja, index=False
            )

        resumen_comparativo.to_excel(writer, sheet_name="resumen_comparativo", index=False)

        # --- Formato, después de escribir todo el contenido ---
        _formatear_hoja_diccionario(
            writer.sheets["00_Diccionario_Variables"], df_diccionario
        )
        for fuente, df in stats_por_fuente.items():
            hoja = fuente if not df.empty else f"{fuente}_vacio"
            if not df.empty:
                _formatear_hoja_estadisticos(writer.sheets[hoja], df)
            else:
                _autoajustar_columnas(writer.sheets[hoja])

        if not resumen_comparativo.empty:
            _autoajustar_columnas(writer.sheets["resumen_comparativo"])
            writer.sheets["resumen_comparativo"].freeze_panes = "A2"